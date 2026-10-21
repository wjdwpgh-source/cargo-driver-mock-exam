"""문제 변환·검사 도구 테스트
사용법: python tests/test_tools.py   (임시 폴더에 복사해서 실행하므로 실제 파일은 바뀌지 않음)
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
errors = []


def check(cond, msg):
    print(("  OK  " if cond else "  FAIL ") + msg)
    if not cond:
        errors.append(msg)


def run(tmp, *args):
    env = dict(os.environ, PYTHONIOENCODING="utf-8")
    p = subprocess.run([sys.executable] + list(args), cwd=tmp, capture_output=True, text=True, encoding="utf-8", env=env)
    return p.returncode, p.stdout + p.stderr


tmp = tempfile.mkdtemp()
for d in ("tools", "data"):
    shutil.copytree(os.path.join(ROOT, d), os.path.join(tmp, d))
os.makedirs(os.path.join(tmp, "images"), exist_ok=True)
sys.path.insert(0, os.path.join(tmp, "tools"))
from openpyxl import load_workbook  # noqa: E402
C_STATUSES = ("draft", "reviewed", "verified", "deprecated")

xlsx = os.path.join(tmp, "data", "raw", "문제_LAW.xlsx")
js = os.path.join(tmp, "data", "questions", "LAW.json")


def load():
    return {q["id"]: q for q in json.load(open(js, encoding="utf-8"))["questions"]}


print("[1] 현재 엑셀 검사·변환")
code, out = run(tmp, "tools/validate.py")
check(code == 0, "엑셀 검사 통과")
code, out = run(tmp, "tools/convert.py")
check(code == 0 and "LAW" in out, "변환 성공")
before = load()
check(len(before) >= 40 and all(q["status"] in C_STATUSES for q in before.values()), "법규 문항 저장 (상태 값 정상)")
orig = {q["id"]: q["version"] for q in json.load(open(os.path.join(ROOT, "data", "questions", "LAW.json"), encoding="utf-8"))["questions"]}
check(all(q["version"] == orig[k] for k, q in before.items()), "변경 없으면 version 유지")

print("[2] 해설 수정 → version +1, 다른 문제는 그대로")
wb = load_workbook(xlsx)
ws = wb["문제"]
head = [c.value for c in ws[1]]
ws.cell(row=2, column=head.index("해설") + 1, value=ws.cell(row=2, column=head.index("해설") + 1).value + " (보완)")
qid2 = ws.cell(row=2, column=1).value
wb.save(xlsx)
run(tmp, "tools/convert.py")
after = load()
check(after[qid2]["version"] == orig[qid2] + 1, "수정한 %s version +1" % qid2)
check(sum(1 for k, q in after.items() if q["version"] == orig[k]) == len(after) - 1, "나머지는 version 그대로")

print("[3] 비고·검수자만 바꾸면 version 그대로")
wb = load_workbook(xlsx)
ws = wb["문제"]
ws.cell(row=3, column=head.index("비고") + 1, value="메모만 수정")
qid3 = ws.cell(row=3, column=1).value
wb.save(xlsx)
run(tmp, "tools/convert.py")
check(load()[qid3]["version"] == orig[qid3], "내용 외 수정은 version 유지")

print("[3-1] 힌트만 바꾸면 version 그대로, JSON에 힌트 반영 (v1.17)")
check(all(q.get("hint") for q in before.values()), "모든 법규 문항에 힌트 있음")
wb = load_workbook(xlsx)
ws = wb["문제"]
ws.cell(row=3, column=head.index("힌트") + 1, value="새 힌트 문구")
wb.save(xlsx)
run(tmp, "tools/convert.py")
check(load()[qid3]["version"] == orig[qid3] and load()[qid3]["hint"] == "새 힌트 문구", "힌트 수정은 version 유지·내용 반영")

print("[3-2] 힌트에 정답 보기 문구 → 경고(변환은 진행)")
wb = load_workbook(xlsx)
ws = wb["문제"]
ans = ws.cell(row=3, column=head.index("정답") + 1).value
ws.cell(row=3, column=head.index("힌트") + 1, value="정답은 " + ws.cell(row=3, column=head.index("보기%d" % ans) + 1).value)
wb.save(xlsx)
code, out = run(tmp, "tools/validate.py")
check(code == 0 and "힌트에 정답 보기 문구" in out, "정답 노출 경고")
ws.cell(row=3, column=head.index("힌트") + 1, value="새 힌트 문구")
wb.save(xlsx)

print("[4] 새 행(id 비움) 추가 → 다음 번호 자동 부여, 엑셀에도 기록")
wb = load_workbook(xlsx)
ws = wb["문제"]
r = len(before) + 2
vals = {"지문": "테스트 문제 지문", "보기1": "가", "보기2": "나", "보기3": "다", "보기4": "라", "정답": 3, "해설": "테스트 해설",
        "관련 법령": "도로교통법 제1조", "기준일": "2026-10-19", "출처": "자체 작성", "난이도": 1, "상태": "draft"}
for k, v in vals.items():
    ws.cell(row=r, column=head.index(k) + 1, value=v)
wb.save(xlsx)
code, out = run(tmp, "tools/convert.py")
new_id = "LAW-%04d" % (len(before) + 1)
check(new_id in load(), "새 번호 %s 부여" % new_id)
check(load_workbook(xlsx)["문제"].cell(row=r, column=1).value == new_id, "엑셀에 번호 기록")

print("[5] 행 삭제 → JSON에서는 deprecated로 남음 (번호 재사용 방지)")
wb = load_workbook(xlsx)
ws = wb["문제"]
ws.delete_rows(r)
wb.save(xlsx)
run(tmp, "tools/convert.py")
check(load()[new_id]["status"] == "deprecated", "지운 문제는 폐기 처리")

print("[6] 오류 검출 → 변환 중단, JSON 그대로")
snapshot = open(js, encoding="utf-8").read()
wb = load_workbook(xlsx)
ws = wb["문제"]
ws.cell(row=4, column=head.index("정답") + 1, value=5)                 # 정답 범위 오류
ws.cell(row=5, column=head.index("보기2") + 1, value=ws.cell(row=5, column=head.index("보기1") + 1).value)  # 보기 중복
ws.cell(row=6, column=head.index("기준일") + 1).value = None          # 법규 기준일 누락 (cell(value=None)은 값을 지우지 않음)
ws.cell(row=7, column=head.index("상태") + 1, value="verified")        # 검수자 없이 verified
ws.cell(row=7, column=head.index("검수자·검수일") + 1).value = None
ws.cell(row=8, column=1, value=ws.cell(row=9, column=1).value)          # id 중복
wb.save(xlsx)
code, out = run(tmp, "tools/convert.py")
check(code == 1 and "중단" in out, "변환 중단")
for word in ("정답은 1~4", "같은 보기", "기준일이 필요", "검수자·검수일이 없음", "id 중복"):
    check(word in out, "오류 메시지: " + word)
check(open(js, encoding="utf-8").read() == snapshot, "JSON 변경 없음")
check("행" in out, "행 번호 안내")

print("[7] 배포 전 JSON 검사 (--json)")
code, out = run(tmp, "tools/validate.py", "--json")
check(code == 0, "JSON 검사 통과")

shutil.rmtree(tmp, ignore_errors=True)

print("[8] 힌트 열이 없는 예전 양식도 변환 / 200자 넘는 힌트는 오류")
tmp = tempfile.mkdtemp()
for d in ("tools", "data"):
    shutil.copytree(os.path.join(ROOT, d), os.path.join(tmp, d))
xlsx = os.path.join(tmp, "data", "raw", "문제_LAW.xlsx")
js = os.path.join(tmp, "data", "questions", "LAW.json")
wb = load_workbook(xlsx)
ws = wb["문제"]
head = [c.value for c in ws[1]]
ws.delete_cols(head.index("힌트") + 1)
wb.save(xlsx)
code, out = run(tmp, "tools/convert.py")
check(code == 0, "힌트 열 없어도 변환 성공")
check(all(not q.get("hint") and q["version"] == orig[k] for k, q in load().items()), "힌트 비움, version 유지")
wb = load_workbook(os.path.join(ROOT, "data", "raw", "문제_LAW.xlsx"))
wb.save(xlsx)
wb = load_workbook(xlsx)
ws = wb["문제"]
head = [c.value for c in ws[1]]
ws.cell(row=2, column=head.index("힌트") + 1, value="가" * 201)
wb.save(xlsx)
code, out = run(tmp, "tools/convert.py")
check(code == 1 and "힌트가 200자를 넘음" in out, "긴 힌트 오류")
wt = load_workbook(os.path.join(ROOT, "data", "raw", "문제_양식.xlsx"))
check("힌트" in [c.value for c in wt["문제"][1]] and wt["문제"].cell(row=1, column=[c.value for c in wt["문제"][1]].index("힌트") + 1).comment is not None,
      "양식에 힌트 열·안내")
shutil.rmtree(tmp, ignore_errors=True)

print("[9] 검수표 만들기 → 판정 입력 → 엑셀 원본 반영 → 변환")
import datetime  # noqa: E402
tmp = tempfile.mkdtemp()
for d in ("tools", "data"):
    shutil.copytree(os.path.join(ROOT, d), os.path.join(tmp, d))
js = os.path.join(tmp, "data", "questions", "LAW.json")
xb = load_workbook(os.path.join(tmp, "data", "raw", "문제_LAW.xlsx"))   # 검수 전 상태로 되돌려 시험
xs = xb["문제"]
hh = [c.value for c in xs[1]]
for r in range(2, xs.max_row + 1):
    if xs.cell(row=r, column=1).value in ("LAW-0001", "LAW-0002", "LAW-0003", "LAW-0004", "LAW-0005"):
        xs.cell(row=r, column=hh.index("상태") + 1).value = "draft"
        xs.cell(row=r, column=hh.index("검수자·검수일") + 1).value = None
xb.save(os.path.join(tmp, "data", "raw", "문제_LAW.xlsx"))
code, out = run(tmp, "tools/make_review_sheet.py", "--all")
sheet = [f for f in os.listdir(os.path.join(tmp, "검수"))][0]
sp = os.path.join(tmp, "검수", sheet)
check(code == 0 and "1차 80 = 80" in out, "검수표 생성 (1차 80문항 표시)")
wb = load_workbook(sp)
ws = wb["검수표"]
h = {c.value: i + 1 for i, c in enumerate(ws[1])}
check(ws.max_row - 1 >= 400 and sum(1 for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=h["1차 80"]).value == "○") == 80, "전체 문항 + 1차 80 표시")
law_rows = [r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=h["과목"]).value == "법규"]
check(any(ws.cell(row=r, column=h["법령 바로가기"]).hyperlink for r in law_rows), "법규 문제에 법령 바로가기 링크")
ids = {ws.cell(row=r, column=h["문제 ID"]).value: r for r in range(2, ws.max_row + 1)}
for qid, v, who in (("LAW-0001", "법령대조완료", "검수자A"), ("LAW-0002", "1차확인", ""), ("LAW-0003", "수정필요", ""), ("LAW-0004", "폐기", ""), ("LAW-0005", "법령대조완료", "")):
    ws.cell(row=ids[qid], column=h["검수 판정"]).value = v
    ws.cell(row=ids[qid], column=h["검수자"]).value = who or None
    ws.cell(row=ids[qid], column=h["검수일"]).value = datetime.date(2026, 10, 21) if who else None
ws.cell(row=ids["LAW-0003"], column=h["수정 의견"]).value = "해설 보완"
wb.save(sp)
code, out = run(tmp, "tools/apply_review.py")
check(code == 0 and "법령대조완료 1" in out and "LAW-0005: 법령대조완료는 검수자·검수일이 필요" in out, "반영 결과 요약·검수자 없는 대조완료는 거부")
run(tmp, "tools/convert.py")
q = {x["id"]: x for x in json.load(open(js, encoding="utf-8"))["questions"]}
check(q["LAW-0001"]["status"] == "verified" and "2026-10-21" in q["LAW-0001"]["law_ref"], "법령대조완료 → verified, 기준일 갱신")
check(q["LAW-0002"]["status"] == "reviewed" and q["LAW-0003"]["status"] == "draft" and q["LAW-0004"]["status"] == "deprecated", "1차확인 → reviewed · 수정필요 → 그대로 · 폐기 → deprecated")
check(q["LAW-0005"]["status"] == "draft", "거부된 문항은 상태 그대로")
xl = load_workbook(os.path.join(tmp, "data", "raw", "문제_LAW.xlsx"))["문제"]
hh = [c.value for c in xl[1]]
note = [xl.cell(row=r, column=hh.index("비고") + 1).value for r in range(2, xl.max_row + 1) if xl.cell(row=r, column=1).value == "LAW-0003"][0]
check("[검수" in note and "해설 보완" in note, "수정필요 의견이 비고에 기록")
shutil.rmtree(tmp, ignore_errors=True)
print("\n결과:", "모두 통과" if not errors else "실패 %d건" % len(errors))
sys.exit(1 if errors else 0)
