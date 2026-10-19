"""문제 입력용 엑셀 양식 만들기 (PRD 8.1)
사용법:  python tools/make_template.py            → data/raw/문제_양식.xlsx (빈 양식)
         python tools/make_template.py LAW        → data/raw/문제_LAW.xlsx (없을 때만 새로 만듦)
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import qcommon as C  # noqa: E402

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
WIDTH = {"id": 11, "지문": 48, "보기1": 22, "보기2": 22, "보기3": 22, "보기4": 22, "정답": 7, "해설": 48, "힌트": 36,
         "관련 법령": 26, "기준일": 12, "출처": 12, "난이도": 8, "이미지": 12, "상태": 11, "검수자·검수일": 18, "비고": 34}
HELP = {
    "id": "비워 두면 변환할 때 자동으로 번호를 붙입니다 (예: LAW-0001). 한 번 붙은 번호는 바꾸지 마세요.",
    "지문": "500자 이내",
    "보기1": "각 150자 이내, 같은 보기 금지",
    "정답": "1~4 숫자",
    "해설": "정답 근거 + 틀린 보기의 핵심 이유",
    "힌트": "선택 입력, 200자 이내. 정답을 직접 말하지 말고 생각의 방향만 알려 주세요. 비워 두면 앱이 틀린 보기 하나를 지워 줍니다.",
    "관련 법령": "법규 과목 필수. 예) 화물자동차 운수사업법 제8조",
    "기준일": "법규 과목 필수. 법령을 확인한 기준 날짜 YYYY-MM-DD",
    "출처": "자체 작성 / 공단 공개 자료 / 이용 허락 자료(허락 근거)",
    "난이도": "1 쉬움 · 2 보통 · 3 어려움",
    "이미지": "images 폴더의 파일 이름 (예: LAW-0001.webp)",
    "상태": "draft 초안 → reviewed 1차 검수 → verified 법령 대조 완료(출제됨) / deprecated 폐기",
    "검수자·검수일": "verified로 바꿀 때 필수. 예) 홍길동 / 2026-10-20",
    "비고": "수정 이력, 근거 확인 메모",
}

HEAD_FILL = PatternFill("solid", start_color="1F5FBF")
INPUT_FILL = PatternFill("solid", start_color="FFFDE7")
THIN = Side(style="thin", color="C9D2DC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build(path, code=None, rows=None, overwrite=False):
    if os.path.exists(path) and not overwrite:
        print("이미 있습니다: %s (덮어쓰지 않음)" % path)
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = C.SHEET
    ws.append(C.HEADERS)
    for i, h in enumerate(C.HEADERS, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        key = h if h in HELP else ("보기1" if h.startswith("보기") else None)
        if key:
            c.comment = Comment(HELP[key], "양식")
        ws.column_dimensions[c.column_letter].width = WIDTH.get(h, 14)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"

    n = max(len(rows or []) + 1, 200)
    for r in range(2, n + 2):
        for i in range(1, len(C.HEADERS) + 1):
            c = ws.cell(row=r, column=i)
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = BORDER
    for ridx, rec in enumerate(rows or [], start=2):
        for i, (h, key) in enumerate(C.COLUMNS, start=1):
            v = rec.get(key, "")
            if key in ("answer", "difficulty") and str(v).isdigit():
                v = int(v)
            ws.cell(row=ridx, column=i, value=v if v != "" else None)

    col = {h: ws.cell(row=1, column=i).column_letter for i, h in enumerate(C.HEADERS, start=1)}
    last = n + 1
    dv_ans = DataValidation(type="whole", operator="between", formula1="1", formula2="4", allow_blank=True,
                            showErrorMessage=True, errorTitle="정답", error="정답은 1~4 숫자로 입력하세요.")
    dv_dif = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True, showErrorMessage=True,
                            errorTitle="난이도", error="1, 2, 3 중에서 고르세요.")
    dv_st = DataValidation(type="list", formula1='"%s"' % ",".join(C.STATUSES), allow_blank=True,
                           showErrorMessage=True, errorTitle="상태", error="목록에서 고르세요.")
    dv_src = DataValidation(type="list", formula1='"자체 작성,공단 공개 자료,이용 허락 자료"', allow_blank=True,
                            showErrorMessage=False)
    for dv, h in ((dv_ans, "정답"), (dv_dif, "난이도"), (dv_st, "상태"), (dv_src, "출처")):
        dv.add("%s2:%s%d" % (col[h], col[h], last))
        ws.add_data_validation(dv)

    # 작성 안내 시트
    g = wb.create_sheet("작성 안내")
    lines = [
        ("문제 입력 양식 – 작성 안내", True),
        ("", False),
        ("1. '문제' 시트에 1행 = 1문항으로 입력합니다. 제목 줄(1행)은 바꾸지 마세요.", False),
        ("2. 노란 칸이 입력 칸입니다. 각 열 제목에 마우스를 올리면 입력 규칙이 보입니다.", False),
        ("3. 상태 흐름: draft(초안) → reviewed(1차 검수: 정답이 하나뿐인지, 보기·해설이 명확한지) → verified(법령 원문 대조 완료).", False),
        ("   verified 문제만 실제 시험에 출제됩니다. 법령 개정 등으로 못 쓰게 된 문제는 지우지 말고 deprecated로 바꾸세요.", False),
        ("4. 법규 과목은 '관련 법령'과 '기준일'(법령을 확인한 날짜)이 꼭 필요합니다. 금액·기간·거리 같은 숫자는 국가법령정보센터 최신 조문과 대조하세요.", False),
        ("5. 시중 교재·유료 사이트 문제를 옮겨 적지 마세요. 출처를 반드시 적습니다(공개 저장소에 올라갑니다).", False),
        ("6. '힌트'는 선택 칸입니다. 정답 보기 문구를 그대로 쓰지 말고 떠올릴 실마리만 적으세요. 힌트를 고쳐도 문제 버전은 바뀌지 않습니다.", False),
        ("7. 입력이 끝나면 폴더의 '문제변환.bat'을 실행하세요. 오류가 있으면 행 번호와 함께 알려 줍니다.", False),
        ("", False),
        ("[예시] 입력 형식 (아래 행은 예시일 뿐 변환되지 않습니다)", True),
    ]
    for text, bold in lines:
        g.append([text])
        g.cell(row=g.max_row, column=1).font = Font(name=FONT, bold=bold, size=12 if bold else 10)
    g.append(C.HEADERS)
    for i in range(1, len(C.HEADERS) + 1):
        c = g.cell(row=g.max_row, column=i)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
    example = ["", "화물자동차 운수사업법상 화물자동차 운수사업에 해당하지 않는 것은?", "화물자동차 운송사업", "화물자동차 운송주선사업",
               "화물자동차 운송가맹사업", "화물자동차 대여사업", 4, "법 제2조: 운수사업은 운송사업·운송주선사업·운송가맹사업을 말한다.",
               "운수사업은 운송·주선·가맹 세 가지입니다. 차를 빌려주는 사업은 다른 법에서 다룹니다.",
               "화물자동차 운수사업법 제2조", "2026-01-01", "자체 작성", 1, "", "draft", "", "[예시]"]
    g.append(example)
    for i in range(1, len(C.HEADERS) + 1):
        g.cell(row=g.max_row, column=i).font = Font(name=FONT, size=10, color="555555")
    g.column_dimensions["A"].width = 14
    for i, h in enumerate(C.HEADERS, start=1):
        if i > 1:
            g.column_dimensions[g.cell(row=1, column=i).column_letter].width = WIDTH.get(h, 14)

    # 입력 칸 표시
    for r in range(2, last + 1):
        for i in range(2, len(C.HEADERS) + 1):
            ws.cell(row=r, column=i).fill = INPUT_FILL
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    print("저장: %s" % path)
    return True


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] in C.CODES:
        build(C.raw_path(sys.argv[1]), sys.argv[1])
    else:
        build(os.path.join(C.RAW, "문제_양식.xlsx"), overwrite=True)
